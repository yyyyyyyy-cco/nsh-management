# -*- coding: utf-8 -*-
"""导出文件的公式注入回归（来源：威胁建模 §十六 / F-47）。

风险：成员**姓名、备注等是用户输入**；openpyxl 会把以 `=` 开头的字符串识别为**公式**
（`data_type='f'`），管理员打开导出的 xlsx 时 Excel 可能求值（可构造 `HYPERLINK`/`DDE` 等
对外请求甚至命令执行）——即 OWASP 所说的 CSV/公式注入一类问题。

修法：导出时对以 `=` `+` `-` `@` 开头的值**显式声明为字符串单元格**（`data_type='s'`），
Excel 便按文本处理。本文件用**往返验证**（写 → 再用 openpyxl 读回断言 `data_type`）证明修复有效。

依赖 openpyxl，缺失时模块级跳过（见 ai-checklist 第 32 条）。
"""
import unittest
from io import BytesIO

try:
    import openpyxl

    from app.utils.excel_export import build_members_xlsx
except ImportError as exc:  # pragma: no cover — 无依赖环境
    import unittest as _unittest

    raise _unittest.SkipTest(f"缺少运行依赖（openpyxl），跳过本模块：{exc}") from exc


class _Member:
    """导出函数只用到这几个属性（轻量替身，避免依赖 ORM）。"""

    def __init__(
        self,
        name: str,
        main_profession: str = "铁衣",
        sub_profession: str | None = None,
        status: str = "active",
        remark: str | None = None,
    ) -> None:
        self.name = name
        self.main_profession = main_profession
        self.sub_profession = sub_profession
        self.status = status
        self.remark = remark


class FormulaInjectionTests(unittest.TestCase):
    """往返验证：写出的单元格必须是**文本**而不是公式。"""

    def _first_data_cell(self, member: _Member, column: int = 1):
        data = build_members_xlsx([member], guild_name="测试帮会", guild_id=1)
        workbook = openpyxl.load_workbook(BytesIO(data))
        sheet = workbook[workbook.sheetnames[0]]
        return sheet.cell(row=3, column=column)

    def test_leading_equals_in_name_is_stored_as_text(self):
        cell = self._first_data_cell(_Member(name='=HYPERLINK("http://evil.example","x")'))
        self.assertEqual(cell.data_type, "s", "以 = 开头的姓名必须以字符串形式写入（不得成为公式）")
        self.assertTrue(str(cell.value).startswith("="), "文本内容本身应原样保留")

    def test_leading_plus_minus_at_are_also_text(self):
        for payload in ("+1+1", "-2+3", "@SUM(A1:A2)"):
            with self.subTest(payload=payload):
                cell = self._first_data_cell(_Member(name=payload))
                self.assertEqual(cell.data_type, "s")

    def test_formula_in_remark_is_also_text(self):
        cell = self._first_data_cell(_Member(name="正常姓名", remark="=cmd|'/C calc'!A1"), column=5)
        self.assertEqual(cell.data_type, "s")

    def test_normal_values_unaffected(self):
        cell = self._first_data_cell(_Member(name="张三"))
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, "张三")

    def test_export_still_produces_readable_workbook(self):
        data = build_members_xlsx([_Member(name="李四")], guild_name="测试帮会", guild_id=1)
        workbook = openpyxl.load_workbook(BytesIO(data))
        self.assertEqual(workbook.sheetnames, ["铁衣"])