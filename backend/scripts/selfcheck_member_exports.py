"""F-2/F-3/F-4 回归：出勤隔离、Excel 来源与新旧回导兼容，仅用内存数据。
运行：backend/.venv/Scripts/python.exe backend/scripts/selfcheck_member_exports.py
"""
import sys
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import unquote

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.v1.members import export_members
from app.core.database import Base
from app.models.attendance import AttendanceRecord
from app.models.guild import Guild
from app.models.member import Member
from app.models.schedule import Schedule
from app.services.member_service import MemberServiceError, attendance_rate
from app.utils.excel_export import HEADERS, build_members_xlsx, member_export_filename
from app.utils.excel_import import ExcelImportError, _parse_workbook, import_members


def sample_members():
    return [Member(id=1, guild_id=1, name="测试甲", main_profession="铁衣", status="formal"),
            Member(id=2, guild_id=1, name="测试乙", main_profession="素问", status="substitute")]


class ExportFormatTests(unittest.TestCase):
    def test_new_export_has_source_and_roundtrips(self):
        blob = build_members_xlsx(sample_members(), "测试帮会甲", 1)
        workbook = load_workbook(BytesIO(blob))
        try:
            self.assertEqual(len(workbook.worksheets), 2)
            for sheet in workbook:
                self.assertEqual((sheet["A1"].value, sheet["B1"].value, sheet["D1"].value), ("所属帮会", "测试帮会甲", 1))
                self.assertEqual([cell.value for cell in sheet[2]], HEADERS)
                self.assertEqual(sheet.freeze_panes, "A3")
        finally:
            workbook.close()
        header, rows = _parse_workbook(blob)
        self.assertEqual(header, ["name", "main_profession", "sub_profession", "status", "remark"])
        self.assertEqual({r[0] for r in rows}, {"测试甲", "测试乙"})

    def test_old_template_still_importable(self):
        workbook = Workbook()
        workbook.active.append(HEADERS)
        workbook.active.append(["旧模板成员", "铁衣", None, "正式", None])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        _, rows = _parse_workbook(stream.getvalue())
        self.assertEqual(rows[0][0], "旧模板成员")

    def test_empty_export_has_valid_sheet(self):
        blob = build_members_xlsx([], "空帮会", 2)
        workbook = load_workbook(BytesIO(blob))
        self.assertEqual(workbook.active["B1"].value, "空帮会")
        self.assertEqual(workbook.active.max_row, 2)
        workbook.close()
        with self.assertRaises(ExcelImportError):
            _parse_workbook(blob)

    def test_guild_source_is_literal_and_filename_safe(self):
        blob = build_members_xlsx(sample_members(), "=1+1", 1)
        workbook = load_workbook(BytesIO(blob), data_only=False)
        self.assertEqual(workbook.active["B1"].data_type, "s")
        workbook.close()
        filename = member_export_filename('甲/乙\\帮会:测\n试', "20260918", "xlsx")
        self.assertEqual(filename, "常驻库_甲_乙_帮会_测_试_20260918.xlsx")


class MemberDataTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        self.session = async_sessionmaker(self.engine, expire_on_commit=False)()
        self.session.add_all([Guild(id=1, name="测试帮会甲"), Guild(id=2, name="测试帮会乙")])
        await self.session.commit()

    async def asyncTearDown(self):
        await self.session.close()
        await self.engine.dispose()

    async def test_attendance_filters_schedule_guild_before_aggregation(self):
        self.session.add_all(sample_members())
        for sid, gid in ((1, 1), (2, 1), (3, 2)):
            self.session.add(Schedule(id=sid, guild_id=gid, opponent="测试", rounds=1, match_time=datetime(2026, 9, 18)))
        await self.session.flush()
        # 第 3 条模拟历史脏引用：同一成员 ID 出现在外帮会赛程，不能污染本帮会出勤率。
        for sid, status in ((1, "normal"), (2, "leave"), (3, "normal")):
            self.session.add(AttendanceRecord(schedule_id=sid, member_id=1, member_name="测试甲",
                                              profession="铁衣", status=status, is_filler=False))
        self.session.add(AttendanceRecord(schedule_id=1, member_id=None, member_name="补人",
                                          profession="铁衣", status="normal", is_filler=True))
        await self.session.commit()
        data = {r["member_id"]: r for r in await attendance_rate(self.session, 1)}
        self.assertEqual(set(data), {1, 2})
        self.assertEqual((data[1]["normal_count"], data[1]["leave_count"], data[1]["attendance_rate"]), (1, 1, 0.5))
        self.assertIsNone(data[2]["attendance_rate"])
        self.assertEqual(await attendance_rate(self.session, 2), [])

    async def test_source_marker_does_not_override_import_guild(self):
        blob = build_members_xlsx(sample_members(), "伪造来源帮会", 999)
        result = await import_members(self.session, 1, blob)
        self.assertEqual(result["imported"], 2)
        rows = list((await self.session.scalars(select(Member))).all())
        self.assertEqual({r.guild_id for r in rows}, {1})

    async def test_unbound_import_rejected_before_parsing(self):
        with self.assertRaises(MemberServiceError) as caught:
            await import_members(self.session, None, b"invalid workbook")
        self.assertEqual(caught.exception.status_code, 403)

    async def test_export_handler_uses_current_guild_and_scoped_records(self):
        self.session.add_all(sample_members())
        self.session.add(Member(guild_id=2, name="外帮会成员", main_profession="铁衣"))
        await self.session.commit()
        response = await export_members(keyword=None, profession=None, status=None, sort_by=None,
                                        sort_order="asc", current_user=SimpleNamespace(guild_id=1), session=self.session)
        self.assertIn("常驻库_测试帮会甲_", unquote(response.headers["content-disposition"]))
        _, rows = _parse_workbook(response.body)
        self.assertEqual({r[0] for r in rows}, {"测试甲", "测试乙"})


if __name__ == "__main__":
    unittest.main(verbosity=2)
