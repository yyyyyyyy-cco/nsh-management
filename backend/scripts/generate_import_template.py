"""生成帮众 Excel 一键导入模板（backend/templates/member_import_template.xlsx）。"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent.parent / "templates" / "member_import_template.xlsx"

HEADERS = ["姓名", "主职业", "副职业", "状态", "备注"]
SAMPLE_ROWS = [
    ["张三", "铁衣", "血河", "正式", "示例：主力坦克"],
    ["李四", "素问", "鸿音", "替补", "示例：替补奶妈"],
    ["王五", "碎梦", "", "正式", ""],
]
WIDTHS = {"A": 14, "B": 12, "C": 12, "D": 10, "E": 28}


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "帮众导入模板"
    ws.append(HEADERS)
    for row in SAMPLE_ROWS:
        ws.append(row)
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(horizontal="center")
    wb.save(OUT)
    print(f"模板已生成：{OUT}")


if __name__ == "__main__":
    main()
