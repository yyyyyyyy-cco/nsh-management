"""Excel 成员导出：按主职业分 Sheet，组内正式在前、替补在后，与导入模板表头一致。"""

from __future__ import annotations

import re
from io import BytesIO
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:  # 仅用于类型标注：本模块是**纯格式化工具**，运行时不需要 ORM（便于独立测试）
    from app.models.member import Member

# 表头与 excel_import 的表头识别一一对应，确保导出文件可直接回导
HEADERS = ["姓名", "主职业", "副职业", "状态", "备注"]
STATUS_LABELS = {"formal": "正式", "substitute": "替补"}
COL_WIDTHS = [14, 14, 14, 10, 30]

# Sheet 名中的 Excel 非法字符（职业名为固定中文，正常不会命中，兜底防御）
_SHEET_ILLEGAL = str.maketrans({c: "-" for c in r":\/?*[]"})

# 会触发 Excel 公式求值的前缀（公式注入，见威胁建模 §十六 / F-47）
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _text_cell(sheet, row: int, column: int, value: object):
    """写入**文本**单元格：对以 `=` `+` `-` `@` 开头的值显式声明为字符串。

    为什么必须显式声明：姓名/备注等是**用户输入**，而 openpyxl 会把以 `=` 开头的字符串识别为
    公式（`data_type='f'`）——管理员打开导出文件时 Excel 可能求值（可构造 `HYPERLINK`/`DDE`
    等对外请求），即 OWASP 归类的 CSV/公式注入。置 `data_type='s'` 后按文本处理，
    内容原样保留（回归用例：`tests/test_excel_export_formula.py` 往返断言）。
    """
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value[:1] in _FORMULA_PREFIXES:
        cell.data_type = "s"
    return cell


def _group_members(members: list[Member], profession_order: list[str] | None = None) -> list[tuple[str, list[Member]]]:
    """按主职业分组：Sheet 顺序按职业目录（profession_order），未知职业兜底到最后；
    组内正式在前、替补在后，同状态按姓名排序。"""
    buckets: dict[str, list[Member]] = {}
    for member in members:
        buckets.setdefault(member.main_profession, []).append(member)

    ordered: list[tuple[str, list[Member]]] = []
    for profession in profession_order or []:
        if profession in buckets:
            ordered.append((profession, buckets.pop(profession)))
    for profession in sorted(buckets):  # 兜底：目录外职业名 / 未提供目录顺序
        ordered.append((profession, buckets[profession]))

    result = []
    for profession, group in ordered:
        group.sort(key=lambda m: (m.status != "formal", m.name))
        result.append((profession, group))
    return result


def member_export_filename(guild_name: str | None, date_tag: str, extension: str) -> str:
    """导出名包含帮会来源；清理路径、控制字符，兼容 Windows 下载。"""
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", guild_name or "未命名帮会")
    return f"常驻库_{name}_{date_tag}.{extension}"


def _write_group(sheet, members: list[Member], guild_name: str | None, guild_id: int | None) -> None:
    # 来源行不是业务表头；导入器仅识别这个固定标记后跳过一行。
    sheet.append(["所属帮会", guild_name or "未命名帮会", "帮会ID", guild_id, "仅作来源标识，非防伪凭证"])
    sheet.cell(row=1, column=2).data_type = "s"
    sheet.cell(row=1, column=1).font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="F3E9D2")
    header_font = Font(bold=True)
    for col, title in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=2, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(col)].width = COL_WIDTHS[col - 1]

    for row, member in enumerate(members, start=3):
        _text_cell(sheet, row, 1, member.name)
        _text_cell(sheet, row, 2, member.main_profession)
        _text_cell(sheet, row, 3, member.sub_profession or "")
        _text_cell(sheet, row, 4, STATUS_LABELS.get(member.status, member.status))
        _text_cell(sheet, row, 5, member.remark or "")

    sheet.freeze_panes = "A3"  # 冻结来源与表头两行


def build_members_xlsx(
    members: list[Member],
    guild_name: str | None = None,
    guild_id: int | None = None,
    profession_order: list[str] | None = None,
) -> bytes:
    """按主职业生成多 Sheet（Sheet 顺序按职业目录），来源标识只用于辨认，不作为导入授权依据。"""
    workbook = Workbook()
    workbook.remove(workbook.active)  # 移除默认空 Sheet

    for profession, group in _group_members(members, profession_order) or [("成员", [])]:
        sheet = workbook.create_sheet(title=profession.translate(_SHEET_ILLEGAL)[:31])
        _write_group(sheet, group, guild_name, guild_id)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
