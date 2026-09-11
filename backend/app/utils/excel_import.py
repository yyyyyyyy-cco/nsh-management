"""Excel 成员导入解析：自动识别表头，重名跳过，返回导入结果。"""
import asyncio
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.member import Member
from app.utils.constants import PROFESSIONS

# 导入限制：与 CSV 导入接口保持一致，防止超大文件/超多行耗尽内存
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB
MAX_IMPORT_ROWS = 5000  # 数据行上限（不含表头）


class ExcelImportError(Exception):
    """导入失败（格式问题），HTTP 400。"""


def _header_index(header: str) -> str:
    """表头归一化：去掉空白，返回字段名。"""
    text = (header or "").strip()
    if not text:
        return ""
    if "姓名" in text or "名字" in text:
        return "name"
    if "主职业" in text or text == "职业":
        return "main_profession"
    if "副职业" in text:
        return "sub_profession"
    if "状态" in text:
        return "status"
    if "备注" in text:
        return "remark"
    return ""


def _parse_status(value) -> str:
    text = str(value or "").strip()
    if text in ("正式", "formal", "1"):
        return "formal"
    if text in ("替补", "substitute", "0"):
        return "substitute"
    return "formal"


def _parse_workbook(content: bytes) -> tuple[list[str], list[tuple]]:
    """同步解析 Excel 工作簿（openpyxl 为 CPU 密集操作，由调用方放入线程池执行）。

    兼容单 Sheet 与导出生成的多 Sheet（按职业分表）文件：遍历所有表头合法的 Sheet。
    返回 (表头, 数据行)；格式问题抛 ExcelImportError。
    """
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ExcelImportError("Excel 文件解析失败，请检查格式")

    # read_only 惰性迭代：边读边计数，超行数上限立即中止，避免恶意文件耗尽内存
    header: list[str] | None = None
    data_rows: list[tuple] = []

    def parse_sheet(sheet) -> bool:
        """解析单个 Sheet，表头合法返回 True；数据行累计受上限保护。"""
        nonlocal header
        sheet_header = None
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index == 0:
                sheet_header = [_header_index(cell) for cell in row]
                if "name" not in sheet_header or "main_profession" not in sheet_header:
                    return False
                if header is None:
                    header = sheet_header
                continue
            if len(data_rows) >= MAX_IMPORT_ROWS:
                raise ExcelImportError(f"成员数据超过 {MAX_IMPORT_ROWS} 行上限，请分批导入")
            data_rows.append(row)
        return True

    try:
        for sheet in workbook.worksheets:
            parse_sheet(sheet)
    finally:
        workbook.close()

    if header is None:
        raise ExcelImportError("表头需包含「姓名」和「职业/主职业」列")
    if not data_rows:
        raise ExcelImportError("Excel 至少需要表头行和一条数据")
    return header, data_rows


async def import_members(session: AsyncSession, guild_id: int, content: bytes) -> dict:
    """解析 Excel 成员数据并入库，返回 {imported, skipped, errors}。重名一律跳过。

    解析（openpyxl）在线程池执行，避免阻塞事件循环；入库逻辑在主线程 session 中执行。
    """
    header, data_rows = await asyncio.to_thread(_parse_workbook, content)

    existing = set(
        (await session.execute(select(Member.name).where(Member.guild_id == guild_id))).scalars()
    )

    imported = 0
    skipped = 0
    errors: list[str] = []
    new_members: list[Member] = []
    for row in data_rows:
        record = {field: (row[index] if index < len(row) else None) for index, field in enumerate(header) if field}
        name = str(record.get("name") or "").strip()
        if not name:
            continue
        # 单元格长度限制：与数据库字段长度对齐，超长计入错误跳过
        if len(name) > 32:
            errors.append(f"{name[:10]}…：姓名超过 32 字符上限")
            skipped += 1
            continue
        if name in existing:
            skipped += 1
            continue
        profession = str(record.get("main_profession") or "").strip()
        if profession not in PROFESSIONS:
            errors.append(f"{name}：无效职业「{profession}」")
            skipped += 1
            continue
        sub_profession = str(record.get("sub_profession") or "").strip() or None
        if sub_profession and len(sub_profession) > 16:
            errors.append(f"{name}：副职业超过 16 字符上限")
            skipped += 1
            continue
        remark = str(record.get("remark") or "").strip() or None
        if remark and len(remark) > 255:
            errors.append(f"{name}：备注超过 255 字符上限")
            skipped += 1
            continue
        new_members.append(
            Member(
                guild_id=guild_id,
                name=name,
                main_profession=profession,
                sub_profession=sub_profession,
                status=_parse_status(record.get("status")),
                remark=remark,
            )
        )
        existing.add(name)

    if new_members:
        try:
            session.add_all(new_members)
            await session.commit()
            imported = len(new_members)
        except Exception:
            await session.rollback()
            raise
    return {"imported": imported, "skipped": skipped, "errors": errors}
